"""Genera docs/Anexo_Implementacion_GeoFire.docx: registro de lo implementado, para pegar o adjuntar al informe.
Uso: pip install python-docx && python scripts/generar_anexo_word.py
"""
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

SALIDA = Path(__file__).resolve().parents[1] / "docs" / "Anexo_Implementacion_GeoFire.docx"
d = Document()
d.styles["Normal"].font.name = "Calibri"
d.styles["Normal"].font.size = Pt(10.5)


def titulo(texto, nivel=1):
    d.add_heading(texto, level=nivel)


def parrafo(texto, negrita=False):
    p = d.add_paragraph()
    p.add_run(texto).bold = negrita
    return p


def lista(items):
    for i in items:
        d.add_paragraph(i, style="List Bullet")


def tabla(cabecera, filas, anchos=None):
    t = d.add_table(rows=1, cols=len(cabecera))
    t.style = "Light Grid Accent 1"
    for i, c in enumerate(cabecera):
        t.rows[0].cells[i].text = c
        for r in t.rows[0].cells[i].paragraphs[0].runs:
            r.bold = True
    for fila in filas:
        celdas = t.add_row().cells
        for i, v in enumerate(fila):
            celdas[i].text = str(v)
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
    d.add_paragraph()


h = d.add_heading("Anexo: registro de lo implementado en GeoFire-Perú", 0)
h.alignment = WD_ALIGN_PARAGRAPH.CENTER
parrafo(f"Proyecto Integrador 2 · UTP · Actualizado al {date.today():%d/%m/%Y}. Repositorio: github.com/Jou3007/Proyect_GeoFire. "
        "Este anexo resume qué se construyó, con qué evidencia y en qué difiere del diseño original.")

titulo("1. Resumen")
tabla(["Bloque", "Estado final"], [
    ["Requisitos funcionales (17)", "15 cumplen, 2 en parte (RF-09 viento como dato y no como capa; RF-13 sin notificación push)"],
    ["Criterios de aceptación (19)", "18 cumplen, 1 en parte (AC-05.2: tiempo de carga con navegador real)"],
    ["Reglas de negocio RN-01 a RN-05", "5 cumplen"],
    ["Casos de prueba CP-01 a CP-19", "19 de 19 pasan, con evidencia y fecha"],
    ["Pruebas automáticas / cobertura", "198 pruebas; cobertura del paquete 79 % (meta 70 %)"],
    ["Pendiente", "Despliegue en la nube; datos oficiales de SERFOR/Cultura; decisión sobre el umbral de NDVI"],
])

titulo("2. Arquitectura y tecnologías")
lista([
    "Lenguaje y web: Python 3.12 y Streamlit (Folium para mapas, Plotly para gráficos).",
    "Datos: PostgreSQL con PostGIS; 20 mil focos de calor históricos, 2.2 mil alertas, 14 distritos y 4 provincias.",
    "Fuentes externas: NASA FIRMS (VIIRS S-NPP, VIIRS NOAA-20, MODIS), Sentinel-2 y Sentinel-1 vía Google Earth Engine, "
    "NASA Worldview/GIBS (humo e imagen VIIRS), OpenStreetMap (asentamientos y aserraderos), Open-Meteo (viento), GADM (distritos).",
    "Ejecución: Docker; ciclo automático cada 3 horas (programador local o GitHub Actions); integración continua con flake8 y pytest.",
    "Seguridad: contraseñas con bcrypt, bloqueo tras 5 intentos, roles, auditoría inmutable, fotos reprocesadas con Pillow.",
])

titulo("3. Qué se implementó por historia de usuario")
tabla(["HU", "Qué se hizo", "Evidencia"], [
    ["HU-01 Entorno", "Repositorio con ramas main/develop protegidas, Docker, PostGIS, pruebas de conexión a Earth Engine y base.",
     "README; test_db.py; test_gee.py"],
    ["HU-02 Ingesta FIRMS", "Descarga VIIRS y MODIS, filtro por la geocerca de Ucayali + 5 km, sin duplicados; 19,864 focos históricos; ciclo cada 3 h.",
     "CP-01 a CP-03; tabla ejecuciones"],
    ["HU-03 NDVI", "NDVI, NDWI y NBR sobre Sentinel-2 con máscara de nubes. Evaluación por zona (distrito y provincia) aunque no haya focos.",
     "CP-04, CP-05; pantalla Zonas en riesgo"],
    ["HU-04 Agua", "Máscara de agua NDWI > −0.1 validada contra radar Sentinel-1: 98.6 % del agua excluida, 0.58 % de tierra excluida por error.",
     "docs/validacion_mascara_agua.md; CP-06"],
    ["HU-05 Visor", "Mapa con filtros por provincia y distrito, línea de tiempo, capas NDVI/estrés/NBR/humo/imagen VIIRS con opacidad y comparación de fechas.",
     "CP-11, CP-12"],
    ["HU-06 Riesgo y alertas", "Motor RN-02, RN-04 y RN-05; alertas; correo con fecha y resultado del envío; reintento si falla; estado No evaluable.",
     "CP-07 a CP-10, CP-17"],
    ["HU-07 Validación", "Pantalla para celular: confirmar o descartar con justificación obligatoria y foto o referencia como evidencia; solo la zona asignada.",
     "CP-15"],
    ["HU-08 Reportes", "PDF, CSV y GeoJSON con los mismos totales, filtros, fuentes, limitaciones y área afectada estimada con dNBR.",
     "CP-16"],
    ["HU-09 Acceso", "Login, tres roles, gestión de usuarios, asignación de zona, pantalla de auditoría y administración de geocercas.",
     "CP-13, CP-14"],
])

titulo("4. Resultados medidos")
tabla(["Medida", "Resultado", "Meta"], [
    ["Ciclo de ingesta con Earth Engine real", "17 s para 177 focos", "≤ 60 s"],
    ["Carga del mapa, solo servidor (20 ejecuciones)", "0.6 a 0.9 s", "< 5 s"],
    ["Carga del mapa, navegador real, 1 usuario", "6.5 s", "< 5 s (no cumple)"],
    ["10 / 25 / 50 usuarios simultáneos", "mapa en 37 s / 68 s / 94 % de errores", "< 5 s (no cumple)"],
    ["Lighthouse, pantalla de ingreso", "Accesibilidad 100; buenas prácticas 100; rendimiento 55", "Accesibilidad ≥ 90"],
    ["Backtesting (499 focos, 2 controles)", "AUC 0.83 frente al azar y 0.68 frente a vecinos", "—"],
    ["Backtesting: umbral NDVI < 0.45", "detecta solo el 1.6 % de los eventos; NDVI < 0.75 detecta el 41.7 %", "—"],
])

titulo("5. Desviaciones respecto al diseño del informe")
lista([
    "Viento: NASA Worldview solo publica viento sobre océanos; se usa Open-Meteo y se muestra como dato en la zona, no como capa.",
    "Límite de Ucayali: se usa FAO GAUL (105,388 km² frente a 102,411 km² oficiales); el administrador puede reemplazarlo con el oficial.",
    "Distritos: GADM 4.1 (2018), 14 distritos; los creados después aparecen dentro de su distrito de origen.",
    "Contraseñas con bcrypt, como indica el informe; las antiguas se migran al iniciar sesión.",
    "Pruebas de carga con Playwright y navegador real, en lugar de Locust (Streamlit usa WebSockets).",
    "RN-02: la cercanía a un asentamiento ya no genera alerta Alta por sí sola (63 % de los focos está a menos de 10 km); agrava otra condición. "
    "Documentado en docs/ajuste_RN-02.md.",
    "Comunidades y aserraderos provienen de OpenStreetMap, no de la base oficial de pueblos indígenas ni de un registro de aserraderos.",
])

titulo("6. Limitaciones conocidas")
lista([
    "Los focos de calor son anomalías térmicas (sobre todo quemas agrícolas), no incendios confirmados por SERFOR.",
    "Los umbrales de riesgo y de zona son provisionales; el backtesting sugiere niveles graduados de NDVI.",
    "Con muchos usuarios simultáneos el visor se satura (Streamlit corre en un solo proceso). Solución: servidor dedicado y varias réplicas.",
    "Sin trabajo sin conexión (RNF-03), tal como anticipa el informe; sin notificaciones push dentro de la interfaz.",
])

titulo("7. Pendiente")
lista([
    "Despliegue gratuito en la nube: base Neon, Streamlit Community Cloud y GitHub Actions (guía en docs/despliegue_nube.md).",
    "Datos oficiales: comunidades nativas (BDPI, Ministerio de Cultura) y lista de incendios confirmados (SERFOR/INDECI).",
    "Decidir con SERFOR el umbral de NDVI y validar el ajuste de RN-02 en la Sprint Review.",
    "Actualizar el informe Word con las desviaciones y resultados de este anexo.",
])

titulo("8. Documentos de respaldo en el repositorio (carpeta docs/)")
lista([
    "revision_informe.md: cumplimiento requisito por requisito.",
    "casos_de_prueba.md: CP-01 a CP-19 con resultado y fecha.",
    "pruebas_rendimiento.md: carga, Lighthouse y causas.",
    "backtesting.md: método, resultados y conclusión.",
    "validacion_mascara_agua.md: método y resultados de la máscara de agua.",
    "ajuste_RN-02.md y despliegue_nube.md.",
])
d.save(SALIDA)
print(f"Escrito {SALIDA}")
